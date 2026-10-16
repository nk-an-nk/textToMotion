#!/usr/bin/env python3
"""Append the translated-input experiments to the preserved original report.

No paid API, GPU or inference is invoked. Uses the original report's drawing
helpers without executing its top-level builder. All 60 source MP4s are embedded.
"""
from __future__ import annotations
import ast
import hashlib
import html
import json
from pathlib import Path
import shutil
import subprocess
import zipfile
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches,Pt

ROOT=Path(__file__).resolve().parents[1]
OLD=ROOT/'reports/presentation-20261016'
OUT=ROOT/'reports/presentation-20261016-updated'
EXP=ROOT/'reports/masked-models-20261016'
if OUT.exists():raise RuntimeError('Existing updated presentation is preserved; move it before rebuilding.')
OUT.mkdir();shutil.copytree(OLD/'assets',OUT/'assets');shutil.copytree(OLD/'media',OUT/'media')
FONT='Noto Sans CJK JP'
NAVY,INK,MUTED='122B3A','183440','5B707B'
TEAL,BLUE,GOLD,RED='007F80','3269AC','B56B13','B84642'
PALE,LINE,WHITE='F0F5F5','DCE6E8','FFFFFF'
prs=Presentation(OLD/'textToMotion_experiment_report_20261016.pptx')
outline=json.loads((OLD/'slide-outline.json').read_text())
embedded_videos=json.loads((OLD/'embedded-video-index.json').read_text())
media=json.loads((OLD/'media-index.json').read_text())
tree=ast.parse((ROOT/'scripts/build_experiment_report.py').read_text())
functions=ast.Module(body=[n for n in tree.body if isinstance(n,ast.FunctionDef)],type_ignores=[])
exec(compile(functions,'report-drawing-helpers','exec'),globals())
summary=read(EXP/'summary.json');request_costs=read(EXP/'request-cost-estimates.json');plan=read(ROOT/'configs/masked_comparison_20261016.json')
LABELS={'molingo':'MoLingo-263','remomask2':'ReMoMask-2','momask':'MoMask'}
P={'lab_proxy':'RTX 4090・12GB制限','b300':'B300'}
S={(s['profile'],s['model']):s for s in summary['summary']}
source='追加実験 2026-10-16 / reports/masked-models-20261016 / seed0の全編を収録'

# Clearly retain the historical conditions instead of relabelling old numbers.
for index,s in enumerate(prs.slides):
    if index:
        for shape in s.shapes:
            if shape.has_text_frame and abs(shape.top/Inches(1)-.35)<.02:
                shape.text_frame.paragraphs[0].runs[0].text='第1期（従来条件） / '+shape.text
        s.notes_slide.notes_text_frame.text+='\n\n更新注記：このページは第1期の条件・数値を保持しています。第2期のLLM英訳・3モデル再測定は31枚目以降を参照。'
        outline[index]['notes']+='\n第1期の条件を保持。第2期の英訳入力再測定とは直接比較しない。'
cover=prs.slides[0]
for shape in cover.shapes:
    if not shape.has_text_frame:continue
    if '7モデルの比較と' in shape.text:
        p=shape.text_frame.paragraphs[0];p.runs[0].text='9モデルの比較と';shape.text_frame.paragraphs[1].runs[0].text='LLM英訳経路の追試'
    if '本編18枚' in shape.text:shape.text_frame.paragraphs[0].runs[0].text='更新版：49枚 / 全60本の動画を埋め込み / 追加実験は31枚目から'
prs.core_properties.title='音声対話 × Text-to-Motion × VRM — 英訳入力の追加実験報告'
prs.core_properties.subject='9モデルの既存検証、MoLingo・ReMoMask-2・MoMaskの2GPU比較、全60動画'

# New posters are sized for slides; full-resolution videos remain unchanged.
for profile in P:
    for model in LABELS:
        for kind in ['dialogue','common']:
            folder=EXP/profile/(kind+'-video')/model
            src=folder/('dialogue.mp4' if kind=='dialogue' else 'combined.mp4')
            ident=f'masked-{profile}-{kind}-{model}'
            shutil.copy2(src,OUT/'media'/f'{ident}.mp4')
            poster=folder/'poster.jpg'
            if not poster.exists():
                when=.55 if kind=='common' else read(folder/'plan.json')['teacher_start_seconds']+1
                subprocess.run(['ffmpeg','-hide_banner','-loglevel','error','-n','-ss',str(when),'-i',str(src),'-frames:v','1',str(poster)],check=True)
            with Image.open(poster) as im:im.resize((960,540),Image.Resampling.LANCZOS).save(OUT/'assets'/f'{ident}.jpg',quality=88,optimize=True)
            media.append({'id':ident,'name':f'{P[profile]} / {LABELS[model]} / {kind}','group':'masked-main','path':f'media/{ident}.mp4','source':str(src.relative_to(ROOT))})
        for i in range(1,6):
            ident=f'masked-{profile}-segment-{model}-{i:02}';src=EXP/profile/'common-video'/model/'segments'/f'{i:02}'/'video.mp4'
            shutil.copy2(src,OUT/'media'/f'{ident}.mp4')
            subprocess.run(['ffmpeg','-hide_banner','-loglevel','error','-n','-ss','0.5','-i',str(src),'-frames:v','1','-vf','scale=960:540','-threads','1','-q:v','3',str(OUT/'assets'/f'{ident}.jpg')],check=True)
            media.append({'id':ident,'name':f'{P[profile]} / {LABELS[model]} / 区間{i}','group':'masked-segment','path':f'media/{ident}.mp4','source':str(src.relative_to(ROOT))})


def movie(s,media_id,x,y,w,h):
    poster=OUT/'assets'/f'{media_id}.jpg';path=OUT/'media'/f'{media_id}.mp4'
    picture(s,poster,x,y,w,h)
    shape=s.shapes.add_movie(str(path),Inches(x),Inches(y),Inches(w),Inches(h),poster_frame_image=str(poster),mime_type='video/mp4')
    embedded_videos.append({'id':media_id,'slide':len(prs.slides),'path':str(path.relative_to(OUT)),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    return shape

# 31 — phase-two findings
s=slide('追加実験：LLM英訳を挟み、3モデルを2環境で実生成',
    '2026-10-16 / MoLingo・ReMoMask-2・MoMask / 共通5区間と実音声の生徒→先生会話',source,
    '第1期は原文直接入力。第2期は利用者の指定に従いLLMで忠実に英訳し、その英語を動作モデルに入力。音声と字幕は原文を保持。共通5区間は全モデルで英訳と音声を共有。会話は実際に受信した教師文を毎回英訳する。B300は今回設定で4090を上回らず、最上位GPUなら必ず高速という結果ではなかった。',section='第2期 / 英訳入力での追加実験')
for x,value,label,note in [( .8,'90回','常駐推論の測定','3モデル×2環境×5文×3 seed'),(4.8,'6回','実音声の会話測定','英訳・生成・通信を含む通し実測'),(8.8,'42本','追加した動画','本編12本＋区間単独30本')]:metric(s,x,2.5,3.72,value,label,note)
callout(s,.8,5.25,11.72,1.5,'主な結果','生成中央値：MoMask 0.168〜0.300秒 / ReMoMask-2 0.277〜0.494秒 / MoLingo 2.217〜3.038秒。',body_size=15)
phase_start=s
jump=text(cover,.8,6.68,6.8,.28,'▶ 今回の追加実験へ（31枚目）',size=12,color='AFC9CE',bold=True);jump.click_action.target_slide=s

# 32 — original/translation pairs and protocol
s=slide('発話内容は保持し、動作モデルへ渡す文だけLLMで英訳',
    '共通5区間は同じ英訳を共有。動作指示への書き換えを行わない。音声・字幕は原文のまま。',source,
    'LLMはOpenAI Realtime gpt-realtime-2.1のテキスト出力。原文・英訳・usage・翻訳時間を保存。第1期の原文直接入力条件とは異なる。T5の日本語未知語問題を回避し、英訳後のMoLingo入力では未知語0件、全3モデルで入力の切り詰めなし。翻訳は事前接続せず、その接続時間を翻訳秒数へ含む。',section='第2期 / 入力条件')
rows=[[u['id'],u['text'],u['motion_prompt'],f"{u['translation']['translation_seconds']:.3f}s"] for u in plan['segments']]
table(s,.76,2.37,[.6,5.0,5.4,.8],['区間','元の発話（音声・字幕）','LLM英訳（動作モデル入力）','英訳'],rows,row_h=.67,font=12)
text(s,.85,6.57,11.6,.31,'会話：受信した教師全文 → LLM英訳 → 常駐モデル → 骨格受信。英訳も通し計測に含める。',size=13,color=TEAL,bold=True)

# 33 — speed table with exact measurement scopes
s=slide('常駐生成はMoMaskが最短。会話の待ちは別に計測',
    '生成中央値：英訳済み共通5区間×3 seed。会話：各条件1回。すべて秒。',source,
    '生成中央値はテキスト符号化・サンプリング・デコード・CPU転送・JSON直列化を含む。モデル読込、共有英訳、通信、保存、MP4描画を除く。会話の骨格受信は生徒の発話終了を基準とし、LLM英訳とJupyter WSS中継を含む。Realtime差は最初の音声チャンク受信との差。会話では英訳表現・音声長が呼び出しごとに変動し、速度順位を確定する反復測定ではない。',section='第2期 / 計測結果')
rows=[]
for profile in P:
 for model in LABELS:
  r=S[profile,model];rows.append([P[profile],LABELS[model],f"{r['generation_median_seconds']:.3f}",f"{r['dialogue_llm_translation_seconds']:.3f}",f"{r['dialogue_bones_after_student_end_seconds']:.3f}",f"+{r['dialogue_gap_from_realtime_seconds']:.3f}"])
table(s,.76,2.4,[2.8,2.0,1.6,1.6,2.0,1.8],['環境','モデル','骨格中央値','会話の英訳','発話終了→骨格','Realtime差'],rows,row_h=.47,font=13)
callout(s,.8,5.97,11.7,.81,'結果の読み方','',body_size=14)
text(s,3.0,6.19,9.1,.30,'今回の設定では、3モデルともRTX 4090の生成時間がB300より短い。',size=15,color=TEAL,bold=True)

# 34 — hardware
s=slide('研究室向けの4090と、前回と同じB300を比較',
    'RTX 5070 / Windowsの実機測定ではない。RAM32GBの強制制限も設けていない。',source,
    '研究室向け代替条件はRTX4090、PyTorch allocator上限12227MiB、CPU8スレッド・affinity8。CUDA context等はallocator上限の外。実CPUはEPYC75F3、割当16vCPU/62GB。B300はXeon6740E、割当24vCPU/251GB、推論8スレッド。両方Ubuntu24.04、Torch2.8.0+cu129、CUDA12.9。同一ソースと重み。動作本体FP32、CLIPは公式FP16、MoLingo T5はFP32、TF32無効、batch1。上半身だけ推論した値ではない。',section='第2期 / ハードウェアと再現条件')
rows=[['GPU','RTX 5070 12GB','RTX 4090 24GB','B300 SXM6 AC・公称288GB'],['CPU','Ryzen 7 5700X','AMD EPYC 75F3','Intel Xeon 6740E'],['CPU制約','8コア16スレッド','推論8 / affinity8','推論8 / affinity8'],['RAM','32GB','割当62GB・上限未設定','割当251GB'],['OS','Windows','Ubuntu 24.04','Ubuntu 24.04'],['GPU時給','研究室機','0.89 USD/h','8.99 USD/h']]
table(s,.76,2.38,[1.6,3.2,3.7,3.3],['項目','研究室実機（未測定）','今回の研究室向け代替','今回の上位GPU'],rows,row_h=.46,font=13)
text(s,.83,5.88,11.65,.77,'全身22関節を生成し、表示時に上半身を使用。腰・脚を固定し、元の骨格も保持。\nTorchのピーク割当：約0.89GiB（MoMask）、1.75〜1.78GiB（ReMoMask-2）、3.10GiB（MoLingo）。',size=16,color=MUTED)

# 35 — cost and retention
s=slide('1リクエストの参考料金を、GPU・音声・英訳に分ける',
    '会話のusageから概算。GPUは生徒音声開始→骨格受信を按分。1ドル200円。ASR・保存等は別。',source,
    'GPU単価は作成時REST API表示の0.89USD/hと8.99USD/h。表は生徒音声開始〜骨格受信のGPU時間按分＋実usageのRealtime音声とLLM英訳。1USD=200円。API単価は第1期資料で2026-10-15確認のgpt-realtime-2.1：100万tokenあたり入力text4/audio32、出力text24/audio64 USD。reasoningはtext内訳として二重加算しない。ASR・起動・待機・保存・動画化を除く参考小計で実請求ではない。今回GPUレンタルはPod作成〜停止で4090約0.343USD、B300約3.942USD、合計約4.286USD。実請求未照合、ストレージ・OpenAI費用別。GPUは停止確認済み。各40GB永続領域に重み・source・venv・結果を24時間保持。期限後削除watcherはこのクラウド環境上のbest effortでprovider TTLではない。',section='第2期 / 費用・保持')
rows=[[P[r['profile']],LABELS[r['model']],f"¥{200*r['gpu_request_window_usd']:.3f}",f"¥{200*r['realtime_voice_estimate_usd']:.3f}",f"¥{200*r['translation_estimate_usd']:.3f}",f"¥{r['request_subtotal_jpy_at_200']:.3f}"] for r in request_costs['rows']]
table(s,.8,2.4,[2.8,2.0,1.6,1.8,1.6,1.9],['環境','モデル','応答中GPU','Realtime音声','LLM英訳','参考小計'],rows,row_h=.42,font=13)
callout(s,.8,5.56,5.65,1.12,'GPUレンタル概算：合計 $4.29','停止済み / 保持・API費用は別',body_size=13)
callout(s,6.75,5.56,5.65,1.12,'保持：10/17 15:37・16:04 JST','停止後24h保持／削除監視はクラウド側。',body_size=12)

# 36 — visual limitations
s=slide('英訳で未知語を回避。発話に合う動作かは別に確認',
    '同じseed0を採用し、見栄えに合わせた選び直しはしない。元の全身骨格を保存。',source,
    'MoLingoのT5で未知語0件。CLIPも入力切り詰めなし。ただし発話を英訳した文は動作記述データセットの入力領域と異なり、自然なco-speech gestureを保証しない。MoMaskの会話では4090約96度、B300約152度の胴体傾斜が発生した。角度は品質スコアではない。FID・自然さ・意味適合の定量評価は実施していない。手指や表情を生成するモデルではなく口パクは音声RMS表示。',section='第2期 / 出力の確認')
for i,model in enumerate(LABELS):
 x=.8+i*4.0;picture(s,OUT/'assets'/f'masked-lab_proxy-dialogue-{model}.jpg',x,2.45,3.72,2.093)
 text(s,x,4.72,3.72,.37,LABELS[model],size=19,bold=True,color=TEAL)
 notes={'molingo':'T5の未知語：0件','remomask2':'一部区間で大きな傾き','momask':'会話出力に大きな体の傾き'}
 text(s,x,5.26,3.72,.58,notes[model],size=15,color=MUTED)
text(s,.86,6.24,11.5,.47,'速度・会話の待ち時間・ジェスチャー品質は分けて評価する。FIDや品質順位は未測定。',size=17,color=GOLD,bold=True)

# 37 — directions and navigation
s=slide('GPUの大型化だけでは、Realtimeとの差1秒に届かない',
    '英訳込み会話の追加待ちは2.59〜6.86秒。ネイティブ推論と待ちの内訳を分けて改善する。',source,
    'この3モデル、batch1、同一精度、Jupyter中継での測定。最初の音声と骨格受信の差は全条件で1秒超。モデル常駐は実施済みで、全文待ち、翻訳接続と生成、通信が残る。今後は短い意味区間での開始と直接通信、翻訳済み文やテキスト特徴の再利用を別条件として測る。研究室5070/Windows実機、連続マイク、複数反復p50/p95、品質評価は未達。次の12枚に全42本の新規動画を埋め込む。',section='第2期 / まとめと全動画')
callout(s,.8,2.4,5.7,2.1,'確認したこと','原文を保持したLLM英訳。\n3モデル×2GPUの実生成。\n90回の常駐測定と会話6回。',body_size=19)
callout(s,6.8,2.4,5.7,2.1,'次に測ること','短い意味区間からの生成開始。\n翻訳・中継・保存の待ちを減らす。\n5070実機と発話への適合性。',body_size=19)
text(s,.86,5.12,11.5,.35,'動画付録：本編12本（38〜43枚目） / 区間単独30本（44〜49枚目）',size=18,color=TEAL,bold=True)
text(s,.86,5.92,11.5,.46,'全60本はPPTXに埋め込み済み。PowerPointで動画をクリックして再生。PDFは静止画。',size=17,color=MUTED)

# 38–43 — both main videos per profile/model
for profile in P:
 for model in LABELS:
  r=S[profile,model]
  s=slide(f'{LABELS[model]} / {"4090・12GB制限" if profile=="lab_proxy" else "B300"}：会話・共通5区間',
      '各動画をクリックして個別に再生。教師の発話とジェスチャーは同じフレームから開始。',source,
      f"英訳経路。会話は実測した待ちを再現するオフライン動画でライブ録画ではない。発話終了→骨格受信{r['dialogue_bones_after_student_end_seconds']:.6f}秒、Realtime差{r['dialogue_gap_from_realtime_seconds']:.6f}秒。共通5区間の常駐生成中央値{r['generation_median_seconds']:.6f}秒。英訳・通信は常駐生成中央値から除く。両動画を同時再生すると音声が重なる。",section='第2期 / 本編動画12本')
  for col,kind in enumerate(['dialogue','common']):
   x=.76+col*6.08;text(s,x,2.2,5.73,.4,'生徒 → 先生の実音声会話' if kind=='dialogue' else '共通の英語学習5区間',size=19,bold=True,color=TEAL)
   ident=f'masked-{profile}-{kind}-{model}';movie(s,ident,x,2.8,5.73,3.223)
   caption=f"骨格受信 {r['dialogue_bones_after_student_end_seconds']:.3f}s / Realtime差 +{r['dialogue_gap_from_realtime_seconds']:.3f}s" if kind=='dialogue' else f"骨格生成中央値 {r['generation_median_seconds']:.3f}s / 英訳時間は別計測"
   text(s,x,6.15,5.73,.40,caption,size=13,color=MUTED)

# 44–49 — every single segment, without omitting any clip
for profile in P:
 for model in LABELS:
  s=slide(f'{LABELS[model]} / {"4090・12GB制限" if profile=="lab_proxy" else "B300"}：区間単独5本',
      '全5区間の動画を埋め込み。1本ずつ再生してください。入力は原文をLLMで英訳した文。',source,
      '共通5区間のseed0動画を切り詰めずに埋め込み。発話・ジェスチャーの開始は区間ごとに同期。元の全身骨格を保存し上半身を表示する。完成動画に合わせたネイティブ骨格の再生速度調整あり。音声は全モデル・環境で共有。',section='第2期 / 区間単独動画30本')
  for i,u in enumerate(plan['segments']):
   row,col=divmod(i,3);x=.78+col*4.13;y=2.28+row*2.29
   movie(s,f'masked-{profile}-segment-{model}-{i+1:02}',x,y,3.86,2.17125)
   text(s,x+.07,y+.03,.7,.22,f'{i+1:02}',size=10,color=WHITE,bold=True)
  callout(s,9.05,4.82,3.88,1.63,'原文の音声・字幕を保持','骨格入力のみLLM英訳。\nフルHD・30fps・上半身表示。',body_size=13)

assert len(prs.slides)==49 and len(embedded_videos)==len(media)==60
pptx_path=OUT/'textToMotion_experiment_report_20261016_updated.pptx'
prs.save(pptx_path)
with zipfile.ZipFile(pptx_path) as z:
 parts=[x for x in z.namelist() if x.startswith('ppt/media/') and x.endswith('.mp4')]
 hashes={hashlib.sha256(z.read(x)).hexdigest() for x in parts}
 assert len(parts)==60 and hashes=={x['sha256'] for x in embedded_videos}
assert pptx_path.stat().st_size<100*1024**2,'PPTX exceeds GitHub regular-file limit'
for name,value in [('embedded-video-index.json',embedded_videos),('slide-outline.json',outline),('media-index.json',media),('measurements.json',{'phase1':read(OLD/'measurements.json'),'phase2':summary,'phase2_cost_estimates':request_costs})]:
 (OUT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n')
notes=['# 報告スライド更新版（49枚・全60動画）','','第1期30枚と18動画を保持。第2期19枚と42動画を追加。','']
for row in outline:notes += [f"## {row['slide']:02}. {row['title']}",'',row['subtitle'],'',row['notes'],'','出典：'+row['source'],'']
(OUT/'speaker-notes.md').write_text('\n'.join(notes))
cards=[]
for m in media:
 poster='assets/'+m['id']+('.jpg' if m['group'].startswith('masked') else '.png')
 cards.append('<article><h3>'+html.escape(m['name'])+'</h3><video controls preload="none" poster="'+poster+'" src="'+m['path']+'"></video><a href="'+m['path']+'">MP4を開く</a></article>')
(OUT/'index.html').write_text('<!doctype html><html lang="ja"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>更新版：全60動画</title><style>body{font-family:system-ui;background:#f0f5f5;color:#183440;padding:30px}main{max-width:1300px;margin:auto}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(340px,1fr));gap:20px}article{background:white;padding:18px}video{width:100%;aspect-ratio:16/9;background:#122b3a}a{color:#3269ac}</style><main><h1>音声対話 × Text-to-Motion × VRM：更新版</h1><p>49枚・全60動画をPowerPointに埋め込み。第1期18本と第2期42本を保持。</p><p><a href="'+pptx_path.name+'">PowerPoint</a> · <a href="'+pptx_path.with_suffix('.pdf').name+'">PDF（静止画）</a> · <a href="speaker-notes.md">ノート</a></p><p>録音済み音声入力による実測を再生するデモです。研究室RTX5070実機は未測定。AvatarSample_VRM1.0 / pixiv VRoid Project。</p><div class="grid">'+''.join(cards)+'</div></main></html>')
(OUT/'README.md').write_text('''# 報告用PowerPoint 更新版（2026-10-16）

49枚、動画60本を埋め込みました。PPTX単体で全動画を保持します。従来30枚・18本は第1期として残し、今回の19枚・42本を追加しています。

- `textToMotion_experiment_report_20261016_updated.pptx`：編集可能な更新版。約96MiB。スライドショーで動画をクリックして再生。
- 同名PDF：閲覧用の静止画。動画は再生できません。
- `index.html` / `media/`：全60本をブラウザー等で個別に再生。
- `speaker-notes.md`：日本語原稿と出典。PPTXのノート欄にも保存。
- `embedded-video-index.json`：全動画のスライド番号とSHA256。

追加実験は31枚目から。31〜37枚目にLLM英訳、測定範囲、90回の推論、会話6回、GPU仕様、料金、出力品質上の課題を整理。38〜43枚目は3モデル×2環境の本編12本、44〜49枚目は全区間単独30本です。

第1期は原文直接入力等の従来条件、第2期はLLM英訳入力です。条件の違う値を直接のモデル順位として比較しないでください。研究室PC実機は未測定で、RTX4090＋allocator12GB上限＋CPU8を代替としています。上半身を表示していますが推論は全身22関節。会話実測は各条件1回、品質やFIDの定量評価は未実施です。MoMaskの一部には大きく体を倒す出力が含まれます。

RunPodは両方停止し、重み・環境・結果を24時間保持しています。期限は10/17 15:37 / 16:04 JST。期限後削除watcherはこのクラウド環境上のbest effortです。

再生成は `python scripts/update_experiment_presentation.py`。既存更新版を上書きしないため、再生成時は出力フォルダーを退避してください。PDFはLibreOfficeから別途出力します。資料作成処理で追加GPUや有料APIを呼びません。元の重み・VRM本体は再配布せず、動画内にアバターのcreditを表示しています。
''')
print(json.dumps({'pptx':str(pptx_path),'slides':len(prs.slides),'embedded_videos':len(parts),'pptx_mib':pptx_path.stat().st_size/1024**2},ensure_ascii=False))
